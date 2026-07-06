import csv
import re
import shutil
from io import StringIO
from pathlib import Path
from urllib.parse import unquote

from openpyxl import load_workbook


SUPPORTED_WIKI_EXTENSIONS = {".md", ".txt", ".csv", ".html", ".htm", ".docx", ".xlsx", ".pdf"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".svg"}


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _csv_to_markdown(text: str) -> str:
    rows = list(csv.reader(StringIO(text)))
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    header = rows[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return "\n".join(lines)


def _docling_image_bytes(value) -> bytes | None:
    if isinstance(value, bytes):
        return value
    if isinstance(value, bytearray | memoryview):
        return bytes(value)
    for attr in ("bytes", "data", "content"):
        data = getattr(value, attr, None)
        if isinstance(data, bytes | bytearray | memoryview):
            return bytes(data)
    return None


def _save_docling_images(document, *, asset_dir: Path, asset_prefix: str, markdown: str) -> str:
    images = getattr(document, "images", None) or {}
    if not hasattr(images, "items"):
        return markdown
    additions: list[str] = []
    asset_dir.mkdir(parents=True, exist_ok=True)
    for raw_name, raw_value in images.items():
        data = _docling_image_bytes(raw_value)
        if not data:
            continue
        name = Path(str(raw_name)).name
        name = re.sub(r"[^0-9A-Za-z._\-\u4e00-\u9fff]+", "-", name).strip("-") or "image.png"
        if not Path(name).suffix:
            name = f"{name}.png"
        (asset_dir / name).write_bytes(data)
        rel = f"{asset_prefix.rstrip('/')}/{name}"
        markdown = markdown.replace(f"]({raw_name})", f"]({rel})").replace(f"]({name})", f"]({rel})")
        if rel not in markdown:
            additions.append(f"![{Path(name).stem}]({rel})")
    return "\n\n".join([part for part in [markdown.strip(), *additions] if part])


def _asset_name(name: str, fallback: str = "image.png") -> str:
    value = re.sub(r"[^0-9A-Za-z._\-\u4e00-\u9fff]+", "-", Path(name).name).strip("-") or fallback
    return value if Path(value).suffix else f"{value}.png"


def _unique_asset_path(asset_dir: Path, name: str) -> Path:
    target = asset_dir / name
    if not target.exists():
        return target
    stem = target.stem
    suffix = target.suffix
    index = 2
    while True:
        candidate = asset_dir / f"{stem}-{index}{suffix}"
        if not candidate.exists():
            return candidate
        index += 1


def _local_image_path(target: str, *, asset_dir: Path, known_assets: dict[str, Path]) -> Path | None:
    value = unquote(target.strip().strip("<>")).replace("\\", "/")
    if not value or re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", value) and not re.match(r"^[a-zA-Z]:/", value):
        return None
    path = Path(value)
    if path.is_file():
        return path
    candidates = [asset_dir / value, asset_dir / Path(value).name]
    candidates.extend(item for name, item in known_assets.items() if name == Path(value).name)
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    return None


def _normalize_docling_asset_links(markdown: str, *, asset_dir: Path, asset_prefix: str) -> str:
    asset_dir = asset_dir.resolve()
    asset_dir.mkdir(parents=True, exist_ok=True)
    known_assets = {item.name: item for item in asset_dir.rglob("*") if item.is_file() and item.suffix.lower() in IMAGE_EXTENSIONS}
    copied: dict[Path, str] = {}

    def replace(match: re.Match) -> str:
        alt, target = match.group(1), match.group(2)
        source = _local_image_path(target, asset_dir=asset_dir, known_assets=known_assets)
        if not source:
            return match.group(0)
        source = source.resolve()
        rel = copied.get(source)
        if not rel:
            if source.parent == asset_dir:
                asset_path = source
            else:
                asset_path = _unique_asset_path(asset_dir, _asset_name(source.name))
                shutil.copyfile(source, asset_path)
            rel = f"{asset_prefix.rstrip('/')}/{asset_path.name}"
            copied[source] = rel
        return f"![{alt}]({rel})"

    markdown = re.sub(r"!\[([^\]]*)\]\(([^)]+)\)", replace, markdown)
    for item in list(asset_dir.iterdir()):
        if item.is_dir():
            shutil.rmtree(item)
    return markdown


def _pdfium_to_markdown(path: Path, *, asset_dir: Path, asset_prefix: str) -> str:
    import pypdfium2 as pdfium

    asset_dir = asset_dir.resolve()
    asset_dir.mkdir(parents=True, exist_ok=True)
    lines: list[str] = [f"# {path.stem}"]
    pdf = pdfium.PdfDocument(str(path))
    try:
        for index in range(len(pdf)):
            page = pdf[index]
            try:
                lines.append(f"## Page {index + 1}")
                text_page = page.get_textpage()
                try:
                    text = text_page.get_text_range().strip()
                finally:
                    text_page.close()
                if text:
                    lines.append(text)
                image_name = f"{path.stem}-page-{index + 1}.png"
                image_path = _unique_asset_path(asset_dir, _asset_name(image_name))
                bitmap = page.render(scale=2)
                try:
                    bitmap.to_pil().save(image_path)
                finally:
                    bitmap.close()
                lines.append(f"![Page {index + 1}]({asset_prefix.rstrip('/')}/{image_path.name})")
            finally:
                page.close()
    finally:
        pdf.close()
    return "\n\n".join(lines)


def _docling_to_markdown(path: Path, *, asset_dir: Path | None = None, asset_prefix: str = "assets") -> str:
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption
    from docling_core.types.doc import ImageRefMode

    pdf_options = PdfPipelineOptions()
    pdf_options.generate_page_images = True
    pdf_options.generate_picture_images = True
    pdf_options.images_scale = 2
    converter = DocumentConverter(format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=pdf_options)})
    try:
        document = converter.convert(path).document
    except Exception:
        if path.suffix.lower() == ".pdf" and asset_dir:
            return _pdfium_to_markdown(path, asset_dir=asset_dir, asset_prefix=asset_prefix)
        raise
    if not asset_dir:
        return document.export_to_markdown()

    asset_dir = asset_dir.resolve()
    asset_dir.mkdir(parents=True, exist_ok=True)
    temp_markdown = asset_dir / f"{path.stem}.md"
    document.save_as_markdown(temp_markdown, artifacts_dir=asset_dir, image_mode=ImageRefMode.REFERENCED)
    markdown = temp_markdown.read_text(encoding="utf-8")
    temp_markdown.unlink(missing_ok=True)
    markdown = _normalize_docling_asset_links(markdown, asset_dir=asset_dir, asset_prefix=asset_prefix)
    return _save_docling_images(document, asset_dir=asset_dir, asset_prefix=asset_prefix, markdown=markdown)


def _xlsx_sheet_headings(path: Path, markdown: str) -> str:
    workbook = load_workbook(path, read_only=True)
    try:
        headings = [f"## {sheet.title}" for sheet in workbook.worksheets if sheet.title not in markdown]
    finally:
        workbook.close()
    return "\n\n".join([*headings, markdown]) if headings else markdown


def convert_to_markdown(path: str | Path, *, asset_dir: str | Path | None = None, asset_prefix: str = "assets") -> str:
    file_path = Path(path)
    ext = file_path.suffix.lower()
    if ext not in SUPPORTED_WIKI_EXTENSIONS:
        raise ValueError(f"unsupported wiki file type: {ext}")
    if ext in {".md", ".txt"}:
        return _read_text(file_path)
    if ext == ".csv":
        return _csv_to_markdown(_read_text(file_path))
    if ext in {".html", ".htm", ".docx", ".xlsx", ".pdf"}:
        image_dir = Path(asset_dir) if asset_dir else None
        markdown = _docling_to_markdown(file_path, asset_dir=image_dir, asset_prefix=asset_prefix)
        if ext == ".xlsx":
            markdown = _xlsx_sheet_headings(file_path, markdown)
        return markdown.strip()
    raise ValueError(f"unsupported wiki file type: {ext}")
