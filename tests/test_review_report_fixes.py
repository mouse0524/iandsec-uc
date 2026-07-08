import io
import importlib.util
import json
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.api.v1.tickets import tickets as ticket_api
from app.controllers import captcha as captcha_module
from app.controllers.user import user_controller
from app.core import dependency as dependency_module
from app.models.enums import TicketActionType, TicketStatus


class Obj:
    def __init__(self, **kwargs):
        self.__dict__.update(kwargs)


class AwaitableQuery:
    def __init__(self, result):
        self.result = result

    def order_by(self, *args):
        return self

    async def count(self):
        return len(self.result)

    async def limit(self, count):
        return self.result[:count]

    async def values_list(self, *args, **kwargs):
        return self.result

    def __await__(self):
        async def _result():
            return self.result

        return _result().__await__()


@pytest.fixture
def anyio_backend():
    return "asyncio"


@pytest.mark.anyio
async def test_captcha_does_not_verify_from_local_cache_when_redis_fails(monkeypatch):
    async def fake_redis(command, *args, **kwargs):
        raise RuntimeError("redis down")

    monkeypatch.setattr(captcha_module, "execute_redis", fake_redis)
    monkeypatch.setattr(captcha_module.CaptchaController, "_generate_code", staticmethod(lambda length=4: "1234"))
    monkeypatch.setattr(captcha_module.CaptchaController, "_generate_image_base64", staticmethod(lambda code: "image"))

    captcha_id, _ = await captcha_module.captcha_controller.create_captcha()

    assert not await captcha_module.captcha_controller.verify_captcha(captcha_id, "1234")


@pytest.mark.anyio
async def test_permission_check_uses_existing_api_cache(monkeypatch):
    async def fake_redis(command, key, *args, **kwargs):
        assert command == "get"
        assert key == "perm:api:user:7:v2"
        return json.dumps(["get/api/v1/ticket/list"], ensure_ascii=False)

    class CachedUser:
        id = 7
        is_superuser = False

        @property
        def roles(self):
            raise AssertionError("permission cache should avoid role/API queries")

    monkeypatch.setattr(dependency_module, "execute_redis", fake_redis, raising=False)

    await dependency_module.PermissionControl.has_permission(
        Obj(method="GET", url=Obj(path="/api/v1/ticket/list")),
        CachedUser(),
    )


@pytest.mark.anyio
async def test_auth_user_does_not_use_noop_auth_cache(monkeypatch):
    calls = []
    user = Obj(
        id=7,
        username="user",
        alias="用户",
        email="user@example.com",
        is_active=True,
        is_superuser=False,
        token_version=0,
    )

    async def fake_redis(command, *args, **kwargs):
        calls.append(command)
        return None

    class UserQuery:
        async def first(self):
            return user

    monkeypatch.setattr("app.controllers.user.execute_redis", fake_redis)
    monkeypatch.setattr(user_controller, "model", Obj(filter=lambda **kwargs: UserQuery()))

    assert await user_controller.get_auth_user(7) is user
    assert calls == []


@pytest.mark.anyio
async def test_tech_related_submitter_ids_uses_cache(monkeypatch):
    async def fake_redis(command, key, *args, **kwargs):
        assert command == "get"
        assert key == "ticket:tech_related_submitters:7:v1"
        return "[21, 22]"

    def fail_dept_query(**kwargs):
        raise AssertionError("cached tech scope should avoid department scan")

    monkeypatch.setattr(ticket_api, "execute_redis", fake_redis)
    monkeypatch.setattr(ticket_api.Dept, "filter", fail_dept_query)

    assert await ticket_api._tech_related_submitter_ids(7) == [21, 22]


@pytest.mark.anyio
async def test_ticket_export_uses_batch_user_query_and_chinese_labels(monkeypatch):
    created = datetime(2026, 7, 7, 12, 0, 0)
    tickets = [
        Obj(
            id=1,
            ticket_no="T-1",
            status=TicketStatus.PENDING_REVIEW,
            project_phase="实施",
            issue_type="现网问题",
            impact_scope="全部",
            category="软件",
            root_cause="",
            title="标题",
            company_name="公司",
            contact_name="联系人",
            email="a@example.com",
            phone="13800000000",
            submitter_id=1,
            reviewer_id=2,
            tech_id=3,
            description="描述",
            reject_reason="",
            created_at=created,
            updated_at=created,
            finished_at=None,
        )
    ]
    actions = [
        Obj(
            ticket_id=1,
            operator_id=1,
            action=TicketActionType.SUBMIT,
            from_status=None,
            to_status=TicketStatus.PENDING_REVIEW,
            comment="提交",
            created_at=created,
        )
    ]
    users = [
        Obj(id=1, alias="张三", username="zhangsan"),
        Obj(id=2, alias="", username="reviewer"),
        Obj(id=3, alias="李四", username="lisi"),
    ]

    async def current_user():
        return Obj(id=99, username="admin", is_superuser=True)

    async def role_names(user):
        return []

    def ticket_filter(*args, **kwargs):
        return AwaitableQuery(tickets)

    def action_filter(**kwargs):
        return AwaitableQuery(actions)

    def attachment_filter(**kwargs):
        return AwaitableQuery([])

    def user_filter(**kwargs):
        assert set(kwargs["id__in"]) == {1, 2, 3}
        return AwaitableQuery(users)

    async def fail_get_user_basic(user_id):
        raise AssertionError("export should batch query users")

    monkeypatch.setattr(ticket_api, "_get_current_user", current_user)
    monkeypatch.setattr(ticket_api, "_get_user_role_names", role_names)
    monkeypatch.setattr(ticket_api.Ticket, "filter", ticket_filter)
    monkeypatch.setattr(ticket_api.TicketAttachment, "filter", attachment_filter)
    monkeypatch.setattr(ticket_api.TicketActionLog, "filter", action_filter)
    monkeypatch.setattr(ticket_api.User, "filter", user_filter)
    monkeypatch.setattr(ticket_api.user_controller, "get_user_basic", fail_get_user_basic)

    response = await ticket_api.export_tickets()
    body = b"".join([chunk async for chunk in response.body_iterator])
    workbook = load_workbook(io.BytesIO(body))
    sheet = workbook.active

    assert sheet.title == "工单"
    assert [cell.value for cell in sheet[1]][:4] == ["工单编号", "状态", "项目阶段", "跟踪"]
    assert sheet["B2"].value == "待客服审核"
    assert sheet["M2"].value == "张三"
    assert "提交" in sheet["S2"].value


@pytest.mark.anyio
async def test_schema_migration_skips_columns_that_already_exist():
    path = (
        Path(__file__).resolve().parents[1]
        / "migrations"
        / "models"
        / "27_20260707130000_move_runtime_schema_fallback_to_migration.py"
    )
    spec = importlib.util.spec_from_file_location("review_report_schema_migration", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    class FakeDB:
        def __init__(self):
            self.columns = {
                ("auditlog", "is_archived"),
                ("ticket", "issue_type"),
                ("project_activity", "activity_type"),
                ("project_activity", "content"),
                ("project_activity", "title"),
            }
            self.tables = {"project_activity"}
            self.scripts = []

        async def execute_query_dict(self, query, values=None):
            if "information_schema.COLUMNS" in query:
                return [{"count": int(tuple(values) in self.columns)}]
            if "information_schema.TABLES" in query:
                return [{"count": int(values[0] in self.tables)}]
            raise AssertionError(query)

        async def execute_script(self, sql):
            self.scripts.append(sql)

    db = FakeDB()

    assert await module.upgrade(db) == "SELECT 1;"
    assert not any("ADD COLUMN `is_archived`" in sql for sql in db.scripts)
    assert not any("ADD COLUMN `issue_type`" in sql for sql in db.scripts)
    assert any("ADD COLUMN `impact_scope`" in sql for sql in db.scripts)
    assert any("CREATE TABLE IF NOT EXISTS `webdav_download_log`" in sql for sql in db.scripts)
